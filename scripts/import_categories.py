"""Import the locally curated Excel rows; never modify Perception or source files.

Run with Python + openpyxl: python scripts/import_categories.py [workspace]
The workbook's explicit inclusion decision and original row order are authoritative.
"""
import hashlib
import json
import re
import shutil
import sys
from pathlib import Path
from urllib.parse import urlparse

import openpyxl

REPO = Path(__file__).resolve().parents[1]
WORKSPACE = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else REPO.parent
SOURCES = [('具身导航', 'navigation', '具身导航文献汇总表_cleaned.xlsx', '具身导航文献汇总表', 20),
           ('具身抓取', 'manipulation', '具身抓取文献汇总表_cleaned.xlsx', '具身抓取文献汇总表', 9),
           ('具身协同', 'collaboration', '具身协同文献汇总表_cleaned.xlsx', '具身协同文献汇总表', 20),
           ('具身规划', 'planning', '具身规划文献.xlsx', '具身规划文献汇总表', 8),
           ('世界模型', 'world-model', '世界模型_元数据最终校对版.xlsx', '世界模型文献汇总表', 13)]
IMAGE_ALIASES = {
    ('manipulation', 5): 'DDPG.png',
    ('world-model', 4): 'NA.png',
    ('world-model', 5): 'MAD (Mapping-Aware Dreamer).png',
    ('world-model', 11): 'Self in Space.png',
}


def text(value):
    return '' if value is None else str(value).strip()


def extract_hyperlink(val):
    if val is None:
        return ''
    val_str = str(val).strip()
    m = re.search(r'HYPERLINK\(\s*["\']([^"\']+)["\']', val_str, re.I)
    if m:
        return m.group(1).strip()
    return val_str


def rows(path, sheet, is_planning=False):
    if is_planning:
        wb = openpyxl.load_workbook(path, data_only=False)
        ws = wb[sheet]
        headers = [c.value for c in ws[1]]
        records = []
        for r_idx in range(2, ws.max_row + 1):
            row_vals = [ws.cell(r_idx, c).value for c in range(1, ws.max_column + 1)]
            if any(v is not None for v in row_vals):
                d = dict(zip(headers, row_vals))
                if '论文pdf' in d and not d.get('论文访问链接'):
                    d['论文访问链接'] = extract_hyperlink(d['论文pdf'])
                if '纳入综述？' in d and '是否纳入综述' not in d:
                    d['是否纳入综述'] = d['纳入综述？']
                if 'github代码链接（如果有的话）' in d and '代码链接（如果有的话）' not in d:
                    d['代码链接（如果有的话）'] = d['github代码链接（如果有的话）']
                if '日期（首次公开）' in d and '日期' not in d:
                    d['日期'] = d['日期（首次公开）']
                records.append((r_idx, d))
        wb.close()
        return records
    else:
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        values = list(book[sheet].values)
        book.close()
        return [(i, dict(zip(values[0], row))) for i, row in enumerate(values[1:], 2)
                if any(value is not None for value in row)]


def split(value):
    return [part.strip() for part in re.split(r'[,，;；\n]+', text(value)) if part.strip()]


def url(value):
    value = text(value)
    if not value or value in ('-', '无', '暂无', 'None'):
        return ''
    if any(k in value for k in ('未发现', '未公开', '待核验', '无', '暂无')):
        return ''
    parsed = urlparse(value)
    assert parsed.scheme in ('https', 'http') and parsed.netloc, value
    return value


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


papers, manifest = [], []
for folder, slug, workbook_file, sheet, expected in SOURCES:
    source = WORKSPACE / folder / workbook_file
    records = rows(source, sheet, is_planning=(slug in ('planning', 'world-model')))
    included = [(i, r) for i, r in records if text(r.get('是否纳入综述', '')).startswith('是') or text(r.get('是否纳入综述', '')).lower().startswith('y')]
    assert len(included) == expected, (source, len(included), expected)
    provenance = {'category': 'embodied-' + slug, 'source': source.relative_to(WORKSPACE).as_posix(),
                  'sha256': sha(source), 'sheet': sheet, 'rawRows': len(records),
                  'includedRows': len(included), 'excludedRows': [i for i, r in records
                  if not (text(r.get('是否纳入综述', '')).startswith('是') or text(r.get('是否纳入综述', '')).lower().startswith('y'))], 'papers': []}
    for row, r in included:
        identity = f'{slug}-{row:03d}'
        image_source = source.parent / text(r['配图'])
        image_note = ''
        if not image_source.is_file():
            image_name = IMAGE_ALIASES.get((slug, row), text(r['方法名称']) + '.png')
            image_source = source.parent / (folder + '文献图片') / image_name
        if slug == 'manipulation' and row == 9:
            supplement = source.parent / (sheet + '_新增AIR-VLA与AIR-VLA+.xlsx')
            match = [v for _, v in rows(supplement, sheet) if v['论文名称'] == r['论文名称']]
            assert len(match) == 1 and match[0]['论文访问链接'] == r['论文访问链接']
            image_source = source.parent / text(match[0]['配图'])
            image_note = supplement.relative_to(WORKSPACE).as_posix()
        image = 'assets/papers/figure-placeholder.svg'
        target = REPO / f'assets/papers/{slug}/{identity}.png'
        retained_image = False
        if image_source.is_file():
            image = f'assets/papers/{slug}/{identity}{image_source.suffix.lower()}'
            target = REPO / image
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(image_source, target)
        elif target.is_file():
            # A source image may be renamed after an earlier verified import.
            # Keep the already imported asset instead of regressing to fallback.
            image = target.relative_to(REPO).as_posix()
            retained_image = True
        if slug == 'world-model' and row == 2 and target.is_file():
            shutil.copyfile(target, REPO / 'assets/covers/world-model-cover.png')
        publication = text(r['发表情况'])
        year_match = re.search(r'\b(?:19|20)\d{2}\b', publication) or re.search(r'\b(?:19|20)\d{2}\b', text(r['日期']))
        year = year_match.group() if year_match else ''
        venue = re.sub(r'\s*\b' + year + r'\b\s*', ' ', publication).strip() if year else publication
        authors = text(r['作者'])
        author_list = [a.strip() for a in authors.split(' and ') if a.strip()] if ' and ' in authors else split(authors)
        paper_item = {'id': identity, 'category': 'embodied-' + slug,
                       'title': text(r['论文名称']), 'methodName': text(r['方法名称']),
                       'authors': author_list, 'taskTags': split(r['任务标签']),
                       'methodTags': split(r['方法标签']), 'datasets': split(r['数据集']),
                       'venue': venue, 'year': year, 'summary': text(r['方法']),
                       'theFirst': text(r['The First']), 'image': image,
                       'paperUrl': url(r.get('论文访问链接', '')), 'codeUrl': url(r.get('代码链接（如果有的话）', ''))}
        if slug == 'world-model':
            paper_item['problem'] = text(r.get('解决问题', ''))
            paper_item['method'] = text(r.get('方法', ''))
            paper_item['datasetText'] = text(r.get('数据介绍', ''))
            paper_item['results'] = text(r.get('实验结果', ''))
        papers.append(paper_item)
        provenance['papers'].append({'id': identity, 'row': row,
            'inclusion': text(r.get('是否纳入综述', '')), 'image': image,
            'imageSource': image_source.relative_to(WORKSPACE).as_posix() if image_source.is_file() else None,
            'imageSha256': sha(target) if target.is_file() else None,
            'imageRetained': retained_image,
            'imageSupplement': image_note or None})
    manifest.append(provenance)

EXPECTED_EXISTING = 49
old_papers = [p for p in papers if p['category'] not in ('embodied-planning', 'embodied-world-model')]
planning_papers = [p for p in papers if p['category'] == 'embodied-planning']
world_model_papers = [p for p in papers if p['category'] == 'embodied-world-model']
assert len(old_papers) == EXPECTED_EXISTING and len({p['id'] for p in old_papers}) == EXPECTED_EXISTING
assert len([p for p in old_papers if p['category'] == 'embodied-navigation']) == 20
assert len([p for p in old_papers if p['category'] == 'embodied-manipulation']) == 9
assert len([p for p in old_papers if p['category'] == 'embodied-collaboration']) == 20
assert len(planning_papers) == 8
assert len(world_model_papers) == 13

EXPECTED_TOTAL = EXPECTED_EXISTING + len(planning_papers) + len(world_model_papers)
assert len(papers) == EXPECTED_TOTAL and len({p['id'] for p in papers}) == EXPECTED_TOTAL

serialized = json.dumps(papers, ensure_ascii=False, indent=2)
(REPO / 'data/category-papers.json').write_text(serialized + '\n', encoding='utf-8')
(REPO / 'assets/js/category-data.js').write_text(
    '// Generated by scripts/import_categories.py; shares the frozen Perception schema.\n'
    'window.PAPERS_DATA.push(...' + serialized + ');\n', encoding='utf-8')
(REPO / 'data/category-provenance.json').write_text(
    json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print(json.dumps({p['category']: p['includedRows'] for p in manifest}, ensure_ascii=False))

